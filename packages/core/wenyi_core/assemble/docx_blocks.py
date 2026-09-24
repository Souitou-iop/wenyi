"""Emit ordered headings, paragraphs, bilingual blocks and basic tables."""

from __future__ import annotations

from typing import Any

from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor

from wenyi_core.assemble.docx_numbering import _list_style_name, _restart_list_numbering
from wenyi_core.assemble.docx_styles import (
    _apply_run_style,
    _fill_paragraph,
    _font_for_text,
    _neutralize_heading_theme_color,
    _segment_style_payload,
    _set_outline_level,
    _style_slices,
)
from wenyi_core.assemble.writer_common import (
    _bilingual_source,
    _seg_text,
)
from wenyi_core.document_styles.docx import (
    text_has_visible_list_prefix,
)
from wenyi_core.ingest.models import KIND_HEADING, Chapter


def _add_heading(
    doc: DocxDocument,
    text: str,
    level: int,
    *,
    style: dict[str, Any] | None = None,
    placements: list[dict[str, Any]] | None = None,
    align: str | None = None,
    shade: str | None = None,
    output_font: str | None = None,
) -> None:
    level = max(1, min(9, level))
    style_name = f"Heading {level}"
    try:
        paragraph = doc.add_heading("", level=level)
    except (KeyError, ValueError):
        paragraph = doc.add_paragraph("")
        try:
            paragraph.style = style_name
        except (KeyError, ValueError):
            pass
    _set_outline_level(paragraph, level)
    _fill_paragraph(
        paragraph,
        text,
        style=style,
        placements=placements,
        align=align,
        shade=shade,
        output_font=output_font,
    )
    # Override the template heading accent color with black when the source has no explicit color.
    _neutralize_heading_theme_color(paragraph, style)


def _add_normal(
    doc: DocxDocument,
    text: str,
    *,
    style: dict[str, Any] | None = None,
    placements: list[dict[str, Any]] | None = None,
    align: str | None = None,
    shade: str | None = None,
    output_font: str | None = None,
    dim: bool = False,
    list_fmt: str | None = None,
    list_ilvl: int = 0,
    list_restart: bool = False,
) -> None:
    if list_fmt is not None:
        style_name = _list_style_name(list_fmt, list_ilvl)
        try:
            paragraph = doc.add_paragraph(style=style_name)
        except KeyError:
            paragraph = doc.add_paragraph()
            style_name = "List Number"
        _fill_paragraph(
            paragraph,
            text,
            style=style,
            placements=placements,
            align=align,
            shade=shade,
            output_font=output_font,
            dim=dim,
        )
        if list_restart:
            _restart_list_numbering(doc, paragraph, style_name=style_name, ilvl=list_ilvl)
        return

    paragraph = doc.add_paragraph()
    _fill_paragraph(
        paragraph,
        text,
        style=style,
        placements=placements,
        align=align,
        shade=shade,
        output_font=output_font,
        dim=dim,
    )


def _add_bilingual_paragraphs(
    doc: DocxDocument,
    source: str,
    target: str,
    order: str,
    *,
    style: dict[str, Any] | None = None,
    placements: list[dict[str, Any]] | None = None,
    align: str | None = None,
    shade: str | None = None,
    output_font: str | None = None,
) -> None:
    src = _bilingual_source(source, target)
    target_font = _font_for_text(output_font, source=source, output_text=target)
    if not src:
        _add_normal(
            doc,
            target,
            style=style,
            placements=placements,
            align=align,
            shade=shade,
            output_font=target_font,
        )
        return
    # Use SimSun on the translated side where applicable; preserve source-side font defaults.
    if order == "source_first":
        _add_normal(doc, src, dim=False, align=align, shade=shade, output_font=None)
        _add_normal(
            doc,
            target,
            style=style,
            placements=placements,
            align=align,
            shade=shade,
            output_font=target_font,
        )
    else:
        _add_normal(
            doc,
            target,
            style=style,
            placements=placements,
            align=align,
            shade=shade,
            output_font=target_font,
        )
        _add_normal(doc, src, dim=True, align=align, shade=shade, output_font=None)


def _flush_table(
    doc: DocxDocument,
    cells: dict[tuple[int, int], tuple[str, str, dict[str, Any]]],
    rows: int,
    cols: int,
    *,
    bilingual: bool,
    order: str,
    output_font: str | None = None,
) -> None:
    table = doc.add_table(rows=rows, cols=cols)
    try:
        table.style = "Table Grid"
    except (KeyError, ValueError):
        pass
    for r in range(rows):
        for c in range(cols):
            target, source, meta = cells.get((r, c), ("", "", {}))
            style, placements, align, shade = _segment_style_payload(
                meta, source=source, output_text=target
            )
            target_font = _font_for_text(output_font, source=source, output_text=target)
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            if bilingual:
                src = _bilingual_source(source, target)
                if src:
                    if order == "source_first":
                        _fill_paragraph(
                            paragraph,
                            src,
                            align=align,
                            shade=shade,
                            output_font=None,
                        )
                        paragraph.add_run("\n")
                        # Place the translation on a second line within the same paragraph.
                        for fragment, frag_style in _style_slices(target, style, placements):
                            run = paragraph.add_run(fragment)
                            _apply_run_style(run, frag_style, output_font=target_font)
                    else:
                        _fill_paragraph(
                            paragraph,
                            target,
                            style=style,
                            placements=placements,
                            align=align,
                            shade=shade,
                            output_font=target_font,
                        )
                        paragraph.add_run("\n")
                        dim_run = paragraph.add_run(src)
                        dim_run.font.size = Pt(9)
                        dim_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                else:
                    _fill_paragraph(
                        paragraph,
                        target,
                        style=style,
                        placements=placements,
                        align=align,
                        shade=shade,
                        output_font=target_font,
                    )
            else:
                _fill_paragraph(
                    paragraph,
                    target,
                    style=style,
                    placements=placements,
                    align=align,
                    shade=shade,
                    output_font=target_font,
                )


def _emit_chapter_blocks(
    doc: DocxDocument,
    chapter: Chapter,
    *,
    bilingual: bool,
    order: str,
    output_font: str | None = None,
) -> None:
    """Write in paragraph order, group contiguous table IDs and merge continuation segments."""
    i = 0
    segs = chapter.segments
    last_list_num_id: int | None = None
    while i < len(segs):
        seg = segs[i]
        meta = seg.meta if isinstance(seg.meta, dict) else {}
        table_id = meta.get("table_id")
        if isinstance(table_id, int):
            cells: dict[tuple[int, int], tuple[str, str, dict[str, Any]]] = {}
            rows = int(meta.get("rows") or 1)
            cols = int(meta.get("cols") or 1)
            while i < len(segs):
                cur = segs[i]
                cur_meta = cur.meta if isinstance(cur.meta, dict) else {}
                if cur_meta.get("table_id") != table_id:
                    break
                r = int(cur_meta.get("row") or 0)
                c = int(cur_meta.get("col") or 0)
                rows = max(rows, int(cur_meta.get("rows") or rows))
                cols = max(cols, int(cur_meta.get("cols") or cols))
                cells[(r, c)] = (_seg_text(cur), cur.source, cur_meta)
                i += 1
            _flush_table(
                doc,
                cells,
                rows,
                cols,
                bilingual=bilingual,
                order=order,
                output_font=output_font,
            )
            last_list_num_id = None
            continue

        if not seg.source.strip() and not (seg.target and seg.target.strip()):
            i += 1
            continue

        target_parts = [_seg_text(seg)]
        source_parts = [seg.source]
        kind = seg.kind
        heading_level = 1
        if kind == KIND_HEADING:
            raw_level = meta.get("heading_level", 1)
            heading_level = raw_level if isinstance(raw_level, int) else 1
        style_meta = meta
        i += 1
        while i < len(segs):
            nxt = segs[i]
            nxt_meta = nxt.meta if isinstance(nxt.meta, dict) else {}
            if not nxt.cont or nxt_meta.get("table_id") is not None:
                break
            target_parts.append(_seg_text(nxt))
            source_parts.append(nxt.source)
            i += 1
        target = "".join(target_parts)
        source = "".join(source_parts)
        style, placements, align, shade = _segment_style_payload(
            style_meta, source=source, output_text=target
        )
        text_font = _font_for_text(output_font, source=source, output_text=target)
        list_num_id = style_meta.get("list_num_id")
        list_ilvl = style_meta.get("list_ilvl")
        list_fmt = style_meta.get("list_fmt")
        is_list = (
            isinstance(list_num_id, int)
            and list_num_id > 0
            and isinstance(list_fmt, str)
            and not text_has_visible_list_prefix(target)
            and not text_has_visible_list_prefix(source)
        )
        if kind == KIND_HEADING:
            _add_heading(
                doc,
                target,
                heading_level,
                style=style,
                placements=placements,
                align=align,
                shade=shade,
                output_font=text_font,
            )
            last_list_num_id = None
        elif bilingual:
            _add_bilingual_paragraphs(
                doc,
                source,
                target,
                order,
                style=style,
                placements=placements,
                align=align,
                shade=shade,
                output_font=output_font,
            )
            last_list_num_id = None
        elif is_list:
            restart = list_num_id != last_list_num_id
            _add_normal(
                doc,
                target,
                style=style,
                placements=placements,
                align=align,
                shade=shade,
                output_font=text_font,
                list_fmt=list_fmt,
                list_ilvl=int(list_ilvl) if isinstance(list_ilvl, int) else 0,
                list_restart=restart,
            )
            last_list_num_id = list_num_id
        else:
            _add_normal(
                doc,
                target,
                style=style,
                placements=placements,
                align=align,
                shade=shade,
                output_font=text_font,
            )
            last_list_num_id = None
